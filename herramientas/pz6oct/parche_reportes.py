"""Uso: python3 parche_reportes.py DIR   (DIR = donde quedaron build_entregas.py, build_pdf_entregas.py,
build_widget_entregas.py y build_xlsx_mo.py escritos desde el skill tablero-entregas-mafesa).
Regla de Moisés (6-oct-2026): en TV, estaciones, PDF, widget y Excel las piezas se muestran como
FALTAN X (en la estación actual) · de TOTAL, y el avance de una estación como 'N listas' (ya pasan a la siguiente).
Idempotente (marca PZ6OCT). Si algo no se encuentra, falla con el texto que faltó: NO publicar reportes a medias."""
import sys,os,re
D=sys.argv[1] if len(sys.argv)>1 else '.'
HELP=open(os.path.join(os.path.dirname(os.path.abspath(__file__)),'pz6.py')).read().split('"""',2)[2]
def patch(fn,pares,inject_after=None):
    p=os.path.join(D,fn)
    if not os.path.exists(p): print('no está (se omite)',fn); return
    t=open(p).read()
    if 'PZ6OCT' in t: print('ya',fn); return
    for a,b in pares:
        n=t.count(a); assert n==1,('%s: NO ENCONTRADO (%d) -> %s'%(fn,n,a[:100])); t=t.replace(a,b)
    if inject_after:
        i=t.index(inject_after)+len(inject_after); t=t[:i]+'\n# PZ6OCT · regla 6-oct-2026 (Moisés): Faltan X de total + N listas\n'+HELP+'\n'+t[i:]
    else: t='# PZ6OCT\n'+t
    open(p,'w').write(t); print('ok',fn)

# 1) datos: plan y entr en cada row (la TV y los reportes los usan)
patch('build_entregas.py',[(
"piezas=int(m['remaining_quantity'] or m['planned_quantity']),\n",
"piezas=int(m['remaining_quantity'] or m['planned_quantity']),plan=int(m['planned_quantity'] or 0),entr=int(m.get('completed_quantity') or 0),\n")])

# 2) PDF
patch('build_pdf_entregas.py',[
 ('("nombre","Nombre",122),("piezas","Piezas",32)','("nombre","Nombre",100),("piezas","Faltan",54)'),
 ("""cx,cw=xs["piezas"]; c.setFont("Helvetica-Bold",7); c.drawCentredString(cx+cw/2,ty,f"{r['piezas']:,}")""",
  """cx,cw=xs["piezas"]; P=calc(r); a_=f"{P['faltan']:,}"; b_=f" de {P['plan']:,}"; wa=pdfmetrics.stringWidth(a_,"Helvetica-Bold",7); wb=pdfmetrics.stringWidth(b_,"Helvetica",5.6); x_=cx+cw/2-(wa+wb)/2
        c.setFillColor(INK); c.setFont("Helvetica-Bold",7); c.drawString(x_,ty,a_); c.setFillColor(MUTED); c.setFont("Helvetica",5.6); c.drawString(x_+wa,ty,b_); c.setFillColor(INK)"""),
 ("""if av and s!="COMPLETED":
                    c.setFillColor(ORANGE)""","""if av and s!="COMPLETED":
                    av=listas_op(r,op) or av; c.setFillColor(ORANGE)"""),
 ("""cx,cw=xs["piezas"]; c.drawCentredString(cx+cw/2,y-rh+4.6,f"{sum(r['piezas'] for r in rows):,}"); y-=rh""",
  """cx,cw=xs["piezas"]; c.drawCentredString(cx+cw/2,y-rh+4.6,f"{sum(calc(r)['faltan'] for r in rows):,} de {sum(calc(r)['plan'] for r in rows):,}"); y-=rh"""),
 ('''c.drawString(x,y-2.2,"115/182"); c.setFillColor(colors.HexColor("#333333")); c.setFont("Helvetica",6.4)
    c.drawString(x+28,y-2.2,"= piezas hechas / total (avance anotado en Katana)"); x+=200''',
  '''c.drawString(x,y-2.2,"85 listas"); c.setFillColor(colors.HexColor("#333333")); c.setFont("Helvetica",6.4)
    c.drawString(x+32,y-2.2,"= hechas en esa estación, pasan a la siguiente · Faltan X de Y = faltan en la estación actual, de las Y de la MO"); x+=360'''),
],inject_after='from reportlab.lib.utils import ImageReader\n')

# 3) widget
patch('build_widget_entregas.py',[
 ('''{(" <b>"+E((r.get("av") or {}).get(o))+"</b>") if (r.get("av") or {}).get(o) and s!="COMPLETED" else ""}''',
  '''{(" <b>"+E(listas_op(r,o) or (r.get("av") or {}).get(o))+"</b>") if (r.get("av") or {}).get(o) and s!="COMPLETED" else ""}'''),
 ('''<span>{r["piezas"]:,} pz</span>''','''<span>Faltan <b>{calc(r)["faltan"]:,}</b> de {calc(r)["plan"]:,}</span>'''),
 ('''<div class="ops">{ops}</div></div>')''','''<div class="ops">{ops}</div>{barra_w(r)}</div>')'''),
 ('''.op b{{color:#92400E}}''','''.op b{{color:#92400E}} .bz{{display:flex;height:7px;border-radius:4px;overflow:hidden;background:#E5E7EB;margin-top:6px}} .bz i{{display:block;height:100%}} .bzl{{display:flex;flex-wrap:wrap;gap:2px 10px;font-size:11.5px;color:#374151;margin-top:3px}} .bzl i{{display:inline-block;width:8px;height:8px;border-radius:2px;margin-right:4px}}'''),
],inject_after='E=html.escape\n')
# barra del widget (después del helper)
p=os.path.join(D,'build_widget_entregas.py'); t=open(p).read() if os.path.exists(p) else 'def barra_w'
if 'def barra_w' not in t:
    t=t.replace('\ndef fd(i):','''
def barra_w(r):
    P=calc(r); T=P['plan'] or 1; w=lambda n:'%.2f%%'%(100*n/T)
    s='<div class="bz"><i style="width:%s;background:#25A35A"></i><i style="width:%s;background:#3E8FD8"></i><i style="width:%s;background:#CBD5E1"></i></div>'%(w(P['entr']),w(P['listas']),w(P['faltan']))
    if P['entr'] or P['listas']:
        s+='<div class="bzl">'+(('<span><i style="background:#25A35A"></i>%s entregadas</span>'%f"{P['entr']:,}") if P['entr'] else '')+(('<span><i style="background:#3E8FD8"></i>%s listas%s</span>'%(f"{P['listas']:,}",(' → '+E(P['sig'])) if P['sig'] else ' para entregar')) if P['listas'] else '')+'<span><i style="background:#CBD5E1"></i>%s faltan%s</span></div>'%(f"{P['faltan']:,}",(' en '+E(P['est'])) if P['est'] else '')
    return s
def fd(i):''',1); open(p,'w').write(t)

# 4) Excel
patch('build_xlsx_mo.py',[
 ("H=['#','MO','# OC','Tipo','SKU','Nombre','Piezas',","H=['#','MO','# OC','Tipo','SKU','Nombre','Faltan (de total)',"),
 ("x['sku'],x['nombre'],x['piezas'],None,","x['sku'],x['nombre'],calc(x)['faltan'],None,"),
 ("if av and s in('IN_PROGRESS','PAUSED'): t='◐ '+av","if av and s in('IN_PROGRESS','PAUSED'): t='◐ '+(listas_op(x,op) or av)"),
 ("ws.cell(r,7).number_format='#,##0';","ws.cell(r,7).number_format='#,##0\" de %s\"'%f\"{calc(x)['plan']:,}\";"),
],inject_after='from openpyxl.drawing.image import Image\n')
print('listo')

# 5) MO en proceso RACKS (skill avance-operaciones): solo si están sus scripts en DIR.
#    Cada fila puede traer entr_mo (completed_quantity) y listas (avance de la estación actual según notas − entr_mo).
#    Qty MO sigue siendo planned_quantity (la identidad del reporte no cambia); se muestra "faltan X de Qty MO".
RK='''
def faltan_rk(r):
    q=int(r.get("qty_mo") or 0); return max(0,q-int(r.get("entr_mo") or 0)-int(r.get("listas") or 0))
'''
if os.path.exists(os.path.join(D,'build_pdf_racks.py')):
    p=os.path.join(D,'build_pdf_racks.py'); t=open(p).read()
    if 'PZ6OCT' not in t:
        for a,b in [('("qty_mo",      "Qty MO",         38)','("qty_mo",      "Faltan de MO",   52)'),
                    ('("descripcion", "Descripción",   114)','("descripcion", "Descripción",   100)'),
                    ('''c.drawCentredString(cx + cw / 2, ty, f"{r['qty_mo']:,}")''',
                     '''a_=f"{faltan_rk(r):,}"; b_=f" de {r['qty_mo']:,}"; wa=pdfmetrics.stringWidth(a_,"Helvetica-Bold",7.4); wb=pdfmetrics.stringWidth(b_,"Helvetica",5.6)
        c.drawString(cx+cw/2-(wa+wb)/2, ty, a_); c.setFont("Helvetica",5.6); c.setFillColor(MUTED); c.drawString(cx+cw/2-(wa+wb)/2+wa, ty, b_); c.setFillColor(INK)'''),
                    ('''f"{sum(r['qty_mo'] for r in rows):,}")\n    return y - row_h''','''f"{sum(faltan_rk(r) for r in rows):,} de {sum(r['qty_mo'] for r in rows):,}")\n    return y - row_h''')]:
            assert t.count(a)==1,('build_pdf_racks.py NO ENCONTRADO',a[:80]); t=t.replace(a,b)
        i=t.index('NUM_KEYS'); t=t[:i]+'# PZ6OCT · regla 6-oct-2026'+RK+t[i:]; open(p,'w').write(t); print('ok build_pdf_racks.py')
if os.path.exists(os.path.join(D,'build_widget_racks.py')):
    p=os.path.join(D,'build_widget_racks.py'); t=open(p).read()
    if 'PZ6OCT' not in t:
        a='''f'<span class="chip"><b>{r["qty_mo"]:,}</b> pzas</span></div>\''''
        b='''f'<span class="chip">Faltan <b>{faltan_rk(r):,}</b> de {r["qty_mo"]:,}</span></div>\''''
        assert t.count(a)==1,'build_widget_racks.py NO ENCONTRADO chip'; t=t.replace(a,b)
        i=t.index('\ndef '); t=t[:i]+'\n# PZ6OCT · regla 6-oct-2026'+RK+t[i:]; open(p,'w').write(t); print('ok build_widget_racks.py')

# ===== 2a parte (Moisés, 6-oct-2026 21:08): "esto mismo en Excel y PDF" — barra + leyenda como en estaciones y TV =====
def patch2(fn,pares,extra=None):
    p=os.path.join(D,fn)
    if not os.path.exists(p): print('no está (se omite)',fn); return
    t=open(p).read()
    if 'PZ6B' in t: print('ya (2a)',fn); return
    assert 'PZ6OCT' in t,'%s: falta la 1a parte del parche'%fn
    for a,b in pares:
        n=t.count(a); assert n==1,('%s: NO ENCONTRADO (%d) -> %s'%(fn,n,a[:100])); t=t.replace(a,b)
    if extra: t=extra(t)
    open(p,'w').write('# PZ6B\n'+t); print('ok (2a)',fn)

def leyenda_txt():
    return '''
def leyenda(r):
    P=calc(r); s=[]
    if P['entr']: s.append(f"{P['entr']:,} entregadas")
    if P['listas']: s.append(f"{P['listas']:,} "+('lista' if P['listas']==1 else 'listas')+((' → '+P['sig']) if P['sig'] else ' para entregar'))
    s.append(f"{P['faltan']:,} faltan"+((' en '+P['est']) if P['est'] else ''))
    return ' · '.join(s)
'''

# PDF: renglón extra de 9 pt solo en MOs con avance: barra (entregadas oscuro, listas medio, faltan claro; se lee en B/N) + leyenda
patch2('build_pdf_entregas.py',[
 ('RH=rh+max(0,len(ocl)-1)*7.2','RH=rh+max(0,len(ocl)-1)*7.2+(10 if (calc(r)["entr"]+calc(r)["listas"])>0 else 0)'),
 ('''        c.setStrokeColor(RULE); c.setLineWidth(.4); c.line(36,y-RH,36+TW,y-RH); y-=RH''',
  '''        P_=calc(r)
        if P_["entr"]+P_["listas"]>0:   # PZ6B barra + leyenda
            yb=y-RH+3.2; x0=xs["sku"][0]+2; lt=leyenda(r); c.setFont("Helvetica",5.6); c.setFillColor(INK); c.drawString(x0,yb,lt)
            bx=x0+pdfmetrics.stringWidth(lt,"Helvetica",5.6)+8; bw=(xs["entrega"][0]-4)-bx; T_=P_["plan"] or 1
            for n_,col_ in ((P_["entr"],"#1F7A3A"),(P_["listas"],"#5B9BD5"),(P_["faltan"],"#E1E5EA")):
                w_=bw*n_/T_
                if w_>0: c.setFillColor(colors.HexColor(col_)); c.rect(bx,yb-0.4,w_,4.2,stroke=0,fill=1); bx+=w_
            c.setFillColor(INK)
        c.setStrokeColor(RULE); c.setLineWidth(.4); c.line(36,y-RH,36+TW,y-RH); y-=RH'''),
 ('''c.drawString(x+32,y-2.2,"= hechas en esa estación, pasan a la siguiente · Faltan X de Y = faltan en la estación actual, de las Y de la MO"); x+=360''',
  '''c.drawString(x+32,y-2.2,"= hechas en esa estación, pasan a la siguiente · Faltan X de Y = faltan en la estación actual, de las Y de la MO"); x+=360
    for col_,l_ in (("#1F7A3A","entregadas"),("#5B9BD5","listas → sig."),("#E1E5EA","faltan")):
        c.setFillColor(colors.HexColor(col_)); c.rect(x,y-3.6,10,5,stroke=0,fill=1); c.setFillColor(colors.HexColor("#333333")); c.setFont("Helvetica",6.4); c.drawString(x+13,y-2.2,l_); x+=pdfmetrics.stringWidth(l_,"Helvetica",6.4)+24'''),
],extra=lambda t:t.replace('\ndef ','\n'+leyenda_txt()+'\ndef ',1))

# Excel: columnas O "Avance" (barra de 20 bloques con colores, como en estaciones) y P "Desglose" (texto); filtro hasta P
patch2('build_xlsx_mo.py',[
 ("'Entrega','Qty en TE']","'Entrega','Qty en TE','Avance','Desglose']"),
 ("ws.auto_filter.ref='A%d:N%d'%(HR,last)","ws.auto_filter.ref='A%d:P%d'%(HR,last)"),
 ("for k,w in zip('ABCDEFGHIJKLMN',[5,16,min(60,max(12,max(OCW or [10])+2)),9,17,30,18,14,14,14,14,16,10,11]): ws.column_dimensions[k].width=w",
  "for k,w in zip('ABCDEFGHIJKLMNOP',[5,16,min(60,max(12,max(OCW or [10])+2)),9,17,30,18,14,14,14,14,16,10,11,24,52]): ws.column_dimensions[k].width=w"),
 ("    ws.cell(r,7).number_format=","    ws.cell(r,15).value=barra_xl(x); ws.cell(r,16,leyenda(x)); ws.cell(r,15).alignment=Alignment(horizontal='left',vertical='center'); ws.cell(r,16).alignment=Alignment(horizontal='left',vertical='center',wrap_text=True)\n    for j_ in (15,16): ws.cell(r,j_).border=Border(top=thin,bottom=thin,left=thin,right=thin)\n    ws.cell(r,7).number_format="),
],extra=lambda t:t.replace('\nH=[','\n'+leyenda_txt()+'''
from openpyxl.cell.rich_text import CellRichText,TextBlock
from openpyxl.cell.text import InlineFont
def barra_xl(x,N=20):
    P=calc(x); T=P['plan'] or 1; e=round(N*P['entr']/T); l=round(N*P['listas']/T); e=min(e,N); l=min(l,N-e); f=N-e-l
    parts=[TextBlock(InlineFont(color=c_,sz=11),'█'*n_) for n_,c_ in ((e,'1F7A3A'),(l,'3E8FD8'),(f,'D9DEE5')) if n_>0]
    return CellRichText(*parts)
H=[''',1))
print('listo 2a')
